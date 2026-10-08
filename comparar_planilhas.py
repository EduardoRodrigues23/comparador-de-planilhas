
import os
import sys
import numbers
import pandas as pd

from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def comparar_planilhas(
    arquivo_a=None,
    arquivo_b=None,
    caminho_relatorio="relatorio_comparacao.xlsx",
    coluna_id_a="ID",
    coluna_id_b="ID"
):
    # Permite executar pelo terminal sem informar os arquivos na função
    if arquivo_a is None:
        arquivo_a = input("Digite o nome da Planilha A: ").strip()

    if arquivo_b is None:
        arquivo_b = input("Digite o nome da Planilha B: ").strip()

    arquivo_a = str(arquivo_a).strip()
    arquivo_b = str(arquivo_b).strip()

    coluna_id_a = str(coluna_id_a).strip()
    coluna_id_b = str(coluna_id_b).strip()

    if not arquivo_a.lower().endswith(".xlsx"):
        arquivo_a += ".xlsx"

    if not arquivo_b.lower().endswith(".xlsx"):
        arquivo_b += ".xlsx"

    if not os.path.isfile(arquivo_a):
        return f"Planilha A não encontrada: {arquivo_a}"

    if not os.path.isfile(arquivo_b):
        return f"Planilha B não encontrada: {arquivo_b}"

    # Carrega as planilhas
    # dtype=str preserva IDs como "001" quando armazenados como texto
    try:
        planilha_a = pd.read_excel(arquivo_a, dtype=str)
    except Exception as erro:
        return f"Não foi possível abrir a Planilha A.\nDetalhes: {erro}"

    try:
        planilha_b = pd.read_excel(arquivo_b, dtype=str)
    except Exception as erro:
        return f"Não foi possível abrir a Planilha B.\nDetalhes: {erro}"

    # Remove espaços dos nomes das colunas
    planilha_a.columns = planilha_a.columns.astype(str).str.strip()
    planilha_b.columns = planilha_b.columns.astype(str).str.strip()

    if planilha_a.columns.duplicated().any():
        return "A Planilha A possui nomes de colunas repetidos."

    if planilha_b.columns.duplicated().any():
        return "A Planilha B possui nomes de colunas repetidos."

    # Verifica se as colunas escolhidas existem
    if coluna_id_a not in planilha_a.columns:
        return (
            f"A Planilha A não possui a coluna "
            f"'{coluna_id_a}'."
        )

    if coluna_id_b not in planilha_b.columns:
        return (
            f"A Planilha B não possui a coluna "
            f"'{coluna_id_b}'."
        )

    # Normaliza os identificadores
    planilha_a[coluna_id_a] = (
        planilha_a[coluna_id_a].astype("string").str.strip()
    )

    planilha_b[coluna_id_b] = (
        planilha_b[coluna_id_b].astype("string").str.strip()
    )

    ids_vazios_a = (
        planilha_a[coluna_id_a].isna()
        | (planilha_a[coluna_id_a] == "")
    ).fillna(True)

    ids_vazios_b = (
        planilha_b[coluna_id_b].isna()
        | (planilha_b[coluna_id_b] == "")
    ).fillna(True)

    if ids_vazios_a.any() or ids_vazios_b.any():
        problemas = []

        if ids_vazios_a.any():
            problemas.append("Planilha A")

        if ids_vazios_b.any():
            problemas.append("Planilha B")

        return (
            "Foram encontrados identificadores vazios em: "
            + ", ".join(problemas)
            + "."
        )

    # Verifica identificadores duplicados
    ids_duplicados_a = (
        planilha_a.loc[
            planilha_a[coluna_id_a].duplicated(),
            coluna_id_a
        ]
        .unique()
        .tolist()
    )

    ids_duplicados_b = (
        planilha_b.loc[
            planilha_b[coluna_id_b].duplicated(),
            coluna_id_b
        ]
        .unique()
        .tolist()
    )

    if ids_duplicados_a or ids_duplicados_b:
        mensagem = "Foram encontrados identificadores duplicados."

        if ids_duplicados_a:
            mensagem += (
                "\nPlanilha A: "
                + ", ".join(map(str, ids_duplicados_a))
            )

        if ids_duplicados_b:
            mensagem += (
                "\nPlanilha B: "
                + ", ".join(map(str, ids_duplicados_b))
            )

        return mensagem

    # Cria uma coluna interna com o mesmo nome nas duas planilhas
    # Assim podemos comparar Pedido com Código, por exemplo
    coluna_interna = "__ID_COMPARADOR__"

    if (
        coluna_interna in planilha_a.columns
        or coluna_interna in planilha_b.columns
    ):
        return (
            "Uma das planilhas possui uma coluna reservada "
            "pelo comparador: __ID_COMPARADOR__."
        )

    planilha_a[coluna_interna] = planilha_a[coluna_id_a]
    planilha_b[coluna_interna] = planilha_b[coluna_id_b]

    # Separa as colunas de dados dos identificadores
    campos_a = [
        coluna for coluna in planilha_a.columns
        if coluna not in (coluna_id_a, coluna_interna)
    ]

    campos_b = [
        coluna for coluna in planilha_b.columns
        if coluna not in (coluna_id_b, coluna_interna)
    ]

    colunas_comuns = [
        coluna for coluna in campos_a
        if coluna in campos_b
    ]

    colunas_so_na_a = [
        coluna for coluna in campos_a
        if coluna not in campos_b
    ]

    colunas_so_na_b = [
        coluna for coluna in campos_b
        if coluna not in campos_a
    ]

    # Se houver apenas os identificadores, não há campos para comparar
    if not colunas_comuns:
        return (
            "As planilhas não possuem colunas de dados em comum "
            "para comparação."
        )

    # Índices permitem localizar cada registro diretamente pelo ID
    dados_a = planilha_a.set_index(coluna_interna, drop=False)
    dados_b = planilha_b.set_index(coluna_interna, drop=False)

    ids_a = pd.Index(planilha_a[coluna_interna])
    ids_b = pd.Index(planilha_b[coluna_interna])

    ids_comuns = ids_a.intersection(ids_b, sort=False)
    ids_so_na_a = ids_a.difference(ids_b, sort=False)
    ids_so_na_b = ids_b.difference(ids_a, sort=False)

    iguais = 0
    diferentes = 0
    detalhes_diferencas = []

    def valores_iguais(valor_a, valor_b):
        if pd.isna(valor_a) and pd.isna(valor_b):
            return True

        if pd.isna(valor_a) or pd.isna(valor_b):
            return False

        return bool(valor_a == valor_b)

    # Compara os registros encontrados nas duas planilhas
    for identificador in ids_comuns:
        linha_a = dados_a.loc[identificador]
        linha_b = dados_b.loc[identificador]

        diferencas_do_registro = []

        for campo in colunas_comuns:
            valor_a = linha_a[campo]
            valor_b = linha_b[campo]

            if valores_iguais(valor_a, valor_b):
                continue

            # Os dados foram lidos como texto para preservar códigos
            # Tentamos calcular diferenças numéricas quando possível
            numero_a = pd.to_numeric(
                pd.Series([valor_a]), errors="coerce"
            ).iloc[0]

            numero_b = pd.to_numeric(
                pd.Series([valor_b]), errors="coerce"
            ).iloc[0]

            if (
                campo.lower() not in ("cliente", "produto", "nome")
                and pd.notna(numero_a)
                and pd.notna(numero_b)
            ):
                diferenca = float(numero_b - numero_a)
            else:
                diferenca = "Alterado"

            diferencas_do_registro.append({
                "ID": identificador,
                "Campo": campo,
                "Planilha A": valor_a,
                "Planilha B": valor_b,
                "Diferença": diferenca
            })

        if diferencas_do_registro:
            diferentes += 1
            detalhes_diferencas.extend(diferencas_do_registro)
        else:
            iguais += 1

    # Registros exclusivos
    registros_so_na_a = planilha_a[
        planilha_a[coluna_interna].isin(ids_so_na_a)
    ].drop(columns=[coluna_interna]).copy()

    registros_so_na_b = planilha_b[
        planilha_b[coluna_interna].isin(ids_so_na_b)
    ].drop(columns=[coluna_interna]).copy()

    relatorio_resumo = pd.DataFrame({
        "Categoria": [
            "Registros iguais",
            "Registros diferentes",
            "Registros só na Planilha A",
            "Registros só na Planilha B"
        ],
        "Quantidade": [
            iguais,
            diferentes,
            len(ids_so_na_a),
            len(ids_so_na_b)
        ]
    })

    relatorio_diferencas = pd.DataFrame(
        detalhes_diferencas,
        columns=[
            "ID",
            "Campo",
            "Planilha A",
            "Planilha B",
            "Diferença"
        ]
    )

    # Gera o relatório Excel
    try:
        with pd.ExcelWriter(
            caminho_relatorio,
            engine="openpyxl"
        ) as escritor:

            relatorio_resumo.to_excel(
                escritor,
                sheet_name="Resumo",
                index=False
            )

            relatorio_diferencas.to_excel(
                escritor,
                sheet_name="Diferenças",
                index=False
            )

            registros_so_na_a.to_excel(
                escritor,
                sheet_name="Só na A",
                index=False
            )

            registros_so_na_b.to_excel(
                escritor,
                sheet_name="Só na B",
                index=False
            )

            for planilha in escritor.book.worksheets:
                planilha.freeze_panes = "A2"
                planilha.auto_filter.ref = planilha.dimensions

                for celula in planilha[1]:
                    celula.font = Font(
                        bold=True,
                        color="FFFFFF"
                    )

                    celula.fill = PatternFill(
                        fill_type="solid",
                        fgColor="1F4E78"
                    )

                    celula.alignment = Alignment(
                        horizontal="center",
                        vertical="center"
                    )

                planilha.row_dimensions[1].height = 25

                for coluna in planilha.columns:
                    maior_tamanho = 0

                    for celula in coluna:
                        if celula.value is not None:
                            tamanho = len(str(celula.value))
                            maior_tamanho = max(
                                maior_tamanho,
                                tamanho
                            )

                    letra = get_column_letter(coluna[0].column)
                    planilha.column_dimensions[letra].width = (
                        min(maior_tamanho + 2, 60)
                    )

        return True

    except PermissionError:
        return (
            "O relatório Excel está aberto ou sem permissão "
            "para gravação. Feche o arquivo e tente novamente."
        )

    except Exception as erro:
        return (
            "Não foi possível gerar o relatório.\n"
            f"Detalhes: {erro}"
        )


if __name__ == "__main__":
    print("================================")
    print("     COMPARADOR DE PLANILHAS")
    print("================================")

    resultado = comparar_planilhas()

    if resultado is True:
        print("Relatório gerado com sucesso!")
    else:
        print("ERRO:", resultado)
